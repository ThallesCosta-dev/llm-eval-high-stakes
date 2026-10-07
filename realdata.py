"""Real-data example (Section 6.7): operational and conduct endpoints from the transfusion-medicine study
(15 models x 15 vignettes x 2 administration conditions, 14 September 2026).
Reads the study's audit package, writes a de-identified long-format CSV to ./data, Figure 7 to ./figures,
and the numbers quoted in the text to ./figures/results_real.json.
"""
import os, io, csv, json, warnings
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from scipy import stats
warnings.filterwarnings("ignore")

HERE = os.path.dirname(os.path.abspath(__file__))
AUDIT = os.path.join(HERE, "..", "Artificial Inteligence", "Audit")
OUT = os.path.join(HERE, "figures"); DATA = os.path.join(HERE, "data")
os.makedirs(OUT, exist_ok=True); os.makedirs(DATA, exist_ok=True)
C1, C2, C3 = "#2a78d6", "#eb6834", "#1baf7a"
INK, INK2, GRID = "#0b0b0b", "#52514e", "#d9d8d4"
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "axes.edgecolor": INK2, "axes.labelcolor": INK,
    "xtick.color": INK2, "ytick.color": INK2, "axes.spines.top": False, "axes.spines.right": False, "axes.grid": True,
    "grid.color": GRID, "grid.linewidth": 0.6, "legend.frameon": False, "figure.dpi": 200, "savefig.dpi": 200,
    "savefig.facecolor": "white"})
rng = np.random.default_rng(1)
se = lambda x: np.std(x, ddof=1)/np.sqrt(len(x))
def boot_ci(x, B=4000):
    x = np.asarray(x, float); idx = rng.integers(0, len(x), (B, len(x)))
    m = x[idx].mean(1); return np.percentile(m, 2.5), np.percentile(m, 97.5)

# ------------------------------------------------------------------ load
import openpyxl
rows = list(csv.DictReader(io.open(os.path.join(AUDIT, "06_supplementary", "S1_per_call_metadata.csv"), encoding="utf-8-sig")))
wb = openpyxl.load_workbook(os.path.join(AUDIT, "02_data", "collection_and_scoring.xlsx"), data_only=True)
ws = wb["Coleta"]; hdr = [c.value for c in ws[1]]
coleta = [dict(zip(hdr, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
col_resp = [h for h in hdr if h.startswith("Resposta")][0]; col_sess = [h for h in hdr if h.startswith("Sess")][0]
wm = wb["Modelos"]; mh = [c.value for c in wm[1]]
modelos = {r[0]: dict(zip(mh, r)) for r in wm.iter_rows(min_row=2, values_only=True) if r[0]}

codes = [f"M{i}" for i in range(1, 16)]
letters = {c: chr(ord("A")+i) for i, c in enumerate(codes)}
def route(c):
    obs = str(modelos[c].get("Observacao") or "")
    return "local" if "local" in obs else "API"
def is_alias(c): return "alias" in str(modelos[c].get("Observacao") or "")
f = lambda v: float(v) if v not in (None, "", "None") else np.nan

# ------------------------------------------------------------------ per-model API (condition S1) summaries
per = {}
long_rows = []
for c in codes:
    rs = [r for r in rows if r["Code"] == c]
    out = np.array([f(r["Output tokens"]) for r in rs]); reas = np.array([f(r["Reasoning tokens"]) for r in rs])
    share = np.array([f(r["Reasoning share (%)"]) for r in rs]); cost = np.array([f(r["Cost upstream (USD)"]) for r in rs])
    ln = np.array([f(r["Answer length (chars)"]) for r in rs]); prov = sorted(set(r["Provider"] for r in rs))
    nonstop = sum(1 for r in rs if r["Finish reason"] != "stop")
    reasoning_exposed = bool(np.nansum(reas) > 0)
    per[c] = dict(n=len(rs), out_med=np.nanmedian(out), share=np.nanmean(share) if reasoning_exposed else np.nan,
                  share_ci=boot_ci(share[~np.isnan(share)]) if reasoning_exposed else (np.nan, np.nan),
                  cost=np.nanmean(cost), cost_ci=boot_ci(cost[~np.isnan(cost)]) if np.nanmax(cost) > 0 else (0, 0),
                  len_med=np.nanmedian(ln), providers=len(prov), alias=is_alias(c), route=route(c), nonstop=nonstop,
                  reasoning=reasoning_exposed, items={r["ResponseID"].split("-")[-1]: dict(out=f(r["Output tokens"]), ln=f(r["Answer length (chars)"])) for r in rs})
    for r in rs:
        long_rows.append(dict(model=letters[c], condition="API, item-by-item", item=r["ResponseID"].split("-")[-1], route=route(c),
            reasoning_exposed=reasoning_exposed, floating_alias=is_alias(c), prompt_tokens=r["Prompt tokens"], output_tokens=r["Output tokens"],
            reasoning_tokens=r["Reasoning tokens"], cost_upstream_usd=r["Cost upstream (USD)"], answer_chars=r["Answer length (chars)"],
            finish_reason=r["Finish reason"], n_providers_for_model=len(prov)))
# chat condition (S2): answer length per item from the workbook text
chat = {}
for r in coleta:
    if str(r[col_sess]) == "2" and r["Modelo"] in codes:
        chat.setdefault(r["Modelo"], {})[r["ItemID"]] = len(str(r[col_resp] or ""))
        long_rows.append(dict(model=letters[r["Modelo"]], condition="Chat, batched", item=r["ItemID"], route=route(r["Modelo"]),
            reasoning_exposed=per[r["Modelo"]]["reasoning"], floating_alias=is_alias(r["Modelo"]), prompt_tokens="", output_tokens="",
            reasoning_tokens="", cost_upstream_usd="", answer_chars=len(str(r[col_resp] or "")), finish_reason="", n_providers_for_model=""))
api_len = {c: {r["ItemID"]: len(str(r[col_resp] or "")) for r in coleta if str(r[col_sess]) == "1" and r["Modelo"] == c} for c in codes}
with io.open(os.path.join(DATA, "study_percall_deidentified.csv"), "w", encoding="utf-8", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=list(long_rows[0].keys())); w.writeheader(); w.writerows(long_rows)

# ------------------------------------------------------------------ paired condition effect on answer length (chars)
paired = {}
for c in codes:
    if c not in chat: continue
    items = sorted(set(api_len[c]) & set(chat[c]))
    d = np.array([chat[c][i] - api_len[c][i] for i in items], float)
    paired[c] = dict(n=len(items), mean=d.mean(), se=se(d), p=stats.ttest_1samp(d, 0).pvalue, median_api=np.median([api_len[c][i] for i in items]), median_chat=np.median([chat[c][i] for i in items]))
pv = np.array([paired[c]["p"] for c in paired]); order = np.argsort(pv); holm = np.empty_like(pv); run = 0.0
for rank, i in enumerate(order):
    run = max(run, min(1.0, (len(pv)-rank)*pv[i])); holm[i] = run
for i, c in enumerate(paired): paired[c]["holm"] = holm[i]
n_shorter = sum(1 for c in paired if paired[c]["mean"] < 0 and paired[c]["holm"] < 0.05)
n_longer = sum(1 for c in paired if paired[c]["mean"] > 0 and paired[c]["holm"] < 0.05)
rel = [paired[c]["median_chat"]/paired[c]["median_api"] for c in paired]

# ------------------------------------------------------------------ shared item effect: pairwise correlation of log output tokens
items_all = [f"I{i:02d}" for i in range(1, 16)]
L = np.full((len(codes), 15), np.nan)
for a, c in enumerate(codes):
    for b, it in enumerate(items_all):
        if it in per[c]["items"] and not np.isnan(per[c]["items"][it]["out"]) and per[c]["items"][it]["out"] > 0:
            L[a, b] = np.log(per[c]["items"][it]["out"])
cors = []
for a in range(len(codes)):
    for b in range(a+1, len(codes)):
        ok = ~np.isnan(L[a]) & ~np.isnan(L[b])
        if ok.sum() >= 8: cors.append(np.corrcoef(L[a, ok], L[b, ok])[0, 1])
cors = np.array(cors)
# two-way decomposition (models x items) of log output tokens on complete cases
okc = ~np.isnan(L).any(0); Lc = L[:, okc]
gm = Lc.mean(); mi = Lc.mean(0); mm = Lc.mean(1)
ss_item = Lc.shape[0]*np.sum((mi-gm)**2); ss_model = Lc.shape[1]*np.sum((mm-gm)**2)
ss_res = np.sum((Lc - mm[:, None] - mi[None, :] + gm)**2); ss_tot = ss_item + ss_model + ss_res

# ------------------------------------------------------------------ numbers for the text
api_costs = {c: per[c]["cost"] for c in codes if per[c]["cost"] > 0}
cmax, cmin = max(api_costs, key=api_costs.get), min(api_costs, key=api_costs.get)
reas_models = [c for c in codes if per[c]["reasoning"]]
res = {
    "R_N_CALLS": str(sum(per[c]["n"] for c in codes)), "R_N_MISSING": str(15*15 - sum(per[c]["n"] for c in codes)),
    "R_N_ALIAS": str(sum(per[c]["alias"] for c in codes)), "R_MAX_PROV": str(max(per[c]["providers"] for c in codes)),
    "R_MAX_PROV_MODEL": letters[max(codes, key=lambda c: per[c]["providers"])],
    "R_COST_MAX": f"{api_costs[cmax]:.3f}", "R_COST_MAX_MODEL": letters[cmax], "R_COST_MIN": f"{api_costs[cmin]:.4f}", "R_COST_MIN_MODEL": letters[cmin],
    "R_COST_RATIO": f"{api_costs[cmax]/api_costs[cmin]:,.0f}",
    "R_COST_TOTAL": f"{sum(np.nansum([f(r['Cost upstream (USD)']) for r in rows if r['Code']==c]) for c in codes):.2f}",
    "R_N_REAS": str(len(reas_models)), "R_SHARE_MIN": f"{min(per[c]['share'] for c in reas_models):.0f}", "R_SHARE_MAX": f"{max(per[c]['share'] for c in reas_models):.0f}",
    "R_SHARE_MIN_MODEL": letters[min(reas_models, key=lambda c: per[c]["share"])], "R_SHARE_MAX_MODEL": letters[max(reas_models, key=lambda c: per[c]["share"])],
    "R_N_CHAT": str(len(paired)), "R_N_SHORTER": str(n_shorter), "R_N_LONGER": str(n_longer),
    "R_REL_MEDIAN": f"{100*np.median(rel):.0f}", "R_REL_MIN": f"{100*min(rel):.0f}", "R_REL_MAX": f"{100*max(rel):.0f}",
    "R_COR_MEAN": f"{cors.mean():.2f}", "R_COR_MIN": f"{cors.min():.2f}", "R_COR_MAX": f"{cors.max():.2f}",
    "R_VAR_ITEM": f"{100*ss_item/ss_tot:.0f}", "R_VAR_MODEL": f"{100*ss_model/ss_tot:.0f}", "R_VAR_RES": f"{100*ss_res/ss_tot:.0f}",
    "R_N_ITEMS_COMPLETE": str(int(okc.sum())), "R_NONSTOP": str(sum(per[c]["nonstop"] for c in codes)),
    "R_VAR_ITEM_WITHIN": f"{100*ss_item/(ss_item+ss_res):.0f}",
}
# Table 6.5 rows
tab = []
for c in codes:
    p = per[c]; ci = p["cost_ci"]
    cost_s = "local (no API cost)" if p["cost"] == 0 else f"{p['cost']:.4f} ({ci[0]:.4f}–{ci[1]:.4f})"
    share_s = "not exposed" if not p["reasoning"] else f"{p['share']:.0f} ({p['share_ci'][0]:.0f}–{p['share_ci'][1]:.0f})"
    chat_s = "—" if c not in paired else f"{paired[c]['mean']/1000:+.1f} ({(paired[c]['mean']-1.96*paired[c]['se'])/1000:+.1f}, {(paired[c]['mean']+1.96*paired[c]['se'])/1000:+.1f}){'*' if paired[c]['holm']<0.05 else ''}"
    tab.append(f"| {letters[c]} | {'yes' if p['alias'] else 'no'} | {p['providers']} | {int(p['out_med']):,} | {share_s} | {cost_s} | {int(p['len_med']):,} | {chat_s} |")
res["R_TABLE"] = "\n".join(tab)
json.dump(res, io.open(os.path.join(OUT, "results_real.json"), "w", encoding="utf-8"), indent=2, ensure_ascii=False)

# ------------------------------------------------------------------ Figure 7
fig, ax = plt.subplots(1, 3, figsize=(7.6, 3.4), gridspec_kw={"width_ratios": [1.1, 1, 1.1]})
ypos = np.arange(len(codes)); lab = [letters[c] for c in codes]
# (a) upstream cost per response, log scale
for i, c in enumerate(codes):
    p = per[c]
    if p["cost"] > 0:
        ax[0].errorbar(p["cost"], i, xerr=[[p["cost"]-p["cost_ci"][0]], [p["cost_ci"][1]-p["cost"]]], fmt="o", color=C1, ecolor=C1, capsize=2, ms=4, lw=1.2)
    else:
        ax[0].text(1.2e-4, i, "local", color=INK2, fontsize=7, va="center")
ax[0].set_xscale("log"); ax[0].set_xlim(1e-4, 1); ax[0].set_yticks(ypos); ax[0].set_yticklabels(lab); ax[0].invert_yaxis()
ax[0].set_xlabel("Upstream cost per response (USD, log)"); ax[0].set_title("(a) Cost", loc="left", fontsize=9, color=INK)
# (b) reasoning share
for i, c in enumerate(codes):
    p = per[c]
    if p["reasoning"]:
        ax[1].errorbar(p["share"], i, xerr=[[p["share"]-p["share_ci"][0]], [p["share_ci"][1]-p["share"]]], fmt="s", color=C2, ecolor=C2, capsize=2, ms=4, lw=1.2)
    else:
        ax[1].text(2, i, "not exposed", color=INK2, fontsize=7, va="center")
ax[1].set_xlim(0, 100); ax[1].set_yticks(ypos); ax[1].set_yticklabels(lab); ax[1].invert_yaxis()
ax[1].set_xlabel("Reasoning tokens (% of output)"); ax[1].set_title("(b) Hidden reasoning", loc="left", fontsize=9, color=INK)
# (c) paired change in answer length, chat minus API
for i, c in enumerate(codes):
    if c in paired:
        m, s = paired[c]["mean"]/1000, 1.96*paired[c]["se"]/1000
        ax[2].errorbar(m, i, xerr=s, fmt="o", color=C3, ecolor=C3, capsize=2, ms=4, lw=1.2)
        if paired[c]["holm"] < 0.05: ax[2].text(m + s + 0.15, i, "*", color=INK, va="center", fontsize=9)
    else:
        ax[2].text(-6.3, i, "API only", color=INK2, fontsize=7, va="center")
ax[2].axvline(0, color=INK2, ls="--", lw=1); ax[2].set_yticks(ypos); ax[2].set_yticklabels(lab); ax[2].invert_yaxis()
ax[2].set_xlabel("Chat − API length (thousand chars)"); ax[2].set_title("(c) Condition effect", loc="left", fontsize=9, color=INK)
ax[2].set_xlim(-6.5, 3)
fig.tight_layout(); fig.savefig(os.path.join(OUT, "fig7_realstudy.png")); plt.close(fig)
print(json.dumps({k: v for k, v in res.items() if k != "R_TABLE"}, indent=1, ensure_ascii=True))
print(res["R_TABLE"].encode("ascii", "replace").decode())
